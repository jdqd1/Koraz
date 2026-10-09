from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUTPUT = r"D:\Jose (Datos)\Medicina\CEDIAH\Web\Plantilla_Carta_de_Trabajo.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["val", "sz", "space", "color"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Arial", size=11, color="222222", bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_placeholder(paragraph, text, size=11):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color="5B6573", bold=True)
    return run


def set_paragraph(paragraph, alignment=None, before=0, after=7, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment


def suppress_paragraph_borders(target):
    element = target._element
    p_pr = element.get_or_add_pPr()
    existing = p_pr.find(qn("w:pBdr"))
    if existing is not None:
        p_pr.remove(existing)
    p_bdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        p_bdr.append(node)
    p_pr.append(p_bdr)


def add_body_paragraph(doc, parts, after=8):
    p = doc.add_paragraph()
    set_paragraph(p, after=after, line=1.18)
    for part in parts:
        if isinstance(part, tuple):
            text, kwargs = part
            run = p.add_run(text)
            set_run_font(run, **kwargs)
        else:
            run = p.add_run(part)
            set_run_font(run)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(34, 34, 34)

title = styles["Title"]
title.font.name = "Arial"
title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
title.font.size = Pt(19)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_before = Pt(4)
title.paragraph_format.space_after = Pt(16)
title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
suppress_paragraph_borders(title)

small = styles.add_style("PlantillaNota", WD_STYLE_TYPE.PARAGRAPH)
small.font.name = "Arial"
small._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
small.font.size = Pt(8.5)
small.font.italic = True
small.font.color.rgb = RGBColor(96, 105, 117)
small.paragraph_format.space_after = Pt(10)
small.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Minimal template instruction, kept outside the letter body.
p = doc.add_paragraph(style="PlantillaNota")
p.add_run("Plantilla editable | Reemplaza todos los campos entre corchetes antes de emitir.")

# Letterhead area.
p = doc.add_paragraph()
set_paragraph(p, after=1, line=1.0)
add_placeholder(p, "[[NOMBRE DE LA EMPRESA]]", size=12)

p = doc.add_paragraph()
set_paragraph(p, after=0, line=1.0)
for text in [
    "RIF/NIT: [[NÚMERO DE IDENTIFICACIÓN FISCAL]]  |  ",
    "Dirección: [[DIRECCIÓN COMPLETA]]",
]:
    run = p.add_run(text)
    set_run_font(run, size=9, color="5B6573")

p = doc.add_paragraph()
set_paragraph(p, after=12, line=1.0)
for text in [
    "Teléfono: [[TELÉFONO]]  |  ",
    "Correo: [[CORREO ELECTRÓNICO]]",
]:
    run = p.add_run(text)
    set_run_font(run, size=9, color="5B6573")

p = doc.add_paragraph()
set_paragraph(p, after=14, line=1.0)
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_placeholder(p, "[[CIUDAD]]", size=10)
run = p.add_run(", ")
set_run_font(run, size=10)
add_placeholder(p, "[[DÍA]]", size=10)
run = p.add_run(" de ")
set_run_font(run, size=10)
add_placeholder(p, "[[MES]]", size=10)
run = p.add_run(" de ")
set_run_font(run, size=10)
add_placeholder(p, "[[AÑO]]", size=10)

p = doc.add_paragraph(style="Title")
pPr = p._p.get_or_add_pPr()
existing_title_borders = pPr.find(qn("w:pBdr"))
if existing_title_borders is not None:
    pPr.remove(existing_title_borders)
p.add_run("Carta de Trabajo")
suppress_paragraph_borders(p)

p = doc.add_paragraph()
set_paragraph(p, after=11, line=1.0)
run = p.add_run("A quien corresponda:")
set_run_font(run, bold=True)

add_body_paragraph(doc, [
    "Por medio de la presente, hacemos constar que el/la señor(a) ",
    ("[[NOMBRE COMPLETO DEL TRABAJADOR]]", {"size": 11, "color": "5B6573", "bold": True}),
    ", titular de ",
    ("[[TIPO Y NÚMERO DE DOCUMENTO DE IDENTIDAD]]", {"size": 11, "color": "5B6573", "bold": True}),
    ", presta sus servicios en ",
    ("[[NOMBRE DE LA EMPRESA]]", {"size": 11, "color": "5B6573", "bold": True}),
    " desde el ",
    ("[[FECHA DE INGRESO]]", {"size": 11, "color": "5B6573", "bold": True}),
    ", desempeñando el cargo de ",
    ("[[CARGO]]", {"size": 11, "color": "5B6573", "bold": True}),
    " en el área de ",
    ("[[ÁREA O DEPARTAMENTO]]", {"size": 11, "color": "5B6573", "bold": True}),
    ".",
])

add_body_paragraph(doc, [
    "Su relación laboral es de carácter ",
    ("[[TIPO DE CONTRATO]]", {"size": 11, "color": "5B6573", "bold": True}),
    " y cumple una jornada de ",
    ("[[JORNADA Y HORARIO]]", {"size": 11, "color": "5B6573", "bold": True}),
    ". Devenga una remuneración de ",
    ("[[SALARIO Y PERIODICIDAD, SI APLICA]]", {"size": 11, "color": "5B6573", "bold": True}),
    ".",
])

add_body_paragraph(doc, [
    "La presente carta se expide a solicitud de la parte interesada, para los fines que estime convenientes.",
], after=17)

p = doc.add_paragraph()
set_paragraph(p, after=18, line=1.0)
run = p.add_run("Atentamente,")
set_run_font(run)

# Signature block without a table so the letter remains easy to edit.
p = doc.add_paragraph()
set_paragraph(p, after=1, line=1.0)
run = p.add_run("\n")
set_run_font(run, size=11)
add_placeholder(p, "________________________________________", size=10)

p = doc.add_paragraph()
set_paragraph(p, after=1, line=1.0)
add_placeholder(p, "[[NOMBRE DE QUIEN FIRMA]]", size=10)

p = doc.add_paragraph()
set_paragraph(p, after=1, line=1.0)
run = p.add_run("[[CARGO DE QUIEN FIRMA]]")
set_run_font(run, size=10, color="5B6573", bold=True)

p = doc.add_paragraph()
set_paragraph(p, after=0, line=1.0)
run = p.add_run("Teléfono: [[TELÉFONO DE CONTACTO]]  |  Correo: [[CORREO DE CONTACTO]]")
set_run_font(run, size=9, color="5B6573")

# Footer with restrained, non-intrusive template label.
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = fp.add_run("Plantilla de carta de trabajo")
set_run_font(fr, size=8, color="8A93A0", italic=True)

doc.core_properties.title = "Plantilla de Carta de Trabajo"
doc.core_properties.subject = "Modelo editable para constancia laboral"
doc.core_properties.author = ""
doc.core_properties.keywords = "carta de trabajo, constancia laboral, plantilla"

doc.save(OUTPUT)
print(OUTPUT)
